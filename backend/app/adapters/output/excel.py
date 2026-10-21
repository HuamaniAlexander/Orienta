"""Lector validado del libro de datos; no interpreta códigos de aprendizaje automático."""

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from app.domain.carrera import Carrera, Fuente

OBLIGATORIOS = {
    "id_carrera",
    "carrera",
    "institucion",
    "duracion_meses",
    "modalidad",
    "ubicacion",
    "costo_mensual_soles",
    "requisitos",
    "campo_laboral",
    "estado_costo",
    "estado_modalidad",
    "fuente_carrera",
    "fuente_costo",
    "fuente_requisitos",
    "fecha_revision",
    "fuente_duracion",
}
UNIVERSIDADES = {"UTP", "UNMSM", "UNI", "UPCH"}


def texto(valor):
    return str(valor).strip() if valor is not None and str(valor).strip() else None


def leer_carreras(ruta: Path) -> list[Carrera]:
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if "Carreras" not in libro.sheetnames:
            raise ValueError("Falta la hoja Carreras")
        filas = iter(libro["Carreras"].values)
        encabezados = next(filas, ())
        if not OBLIGATORIOS.issubset(encabezados) or len(encabezados) != len(set(encabezados)):
            raise ValueError("Columnas obligatorias ausentes o duplicadas")
        resultado, identificadores, identidades = [], set(), set()
        for numero_fila, valores in enumerate(filas, 2):
            if all(v is None for v in valores):
                continue
            fila = dict(zip(encabezados, valores))
            try:

                def numerico(clave):
                    valor = texto(fila[clave])
                    return Decimal(valor) if valor is not None else None

                duracion = numerico("duracion_meses")
                if duracion is not None and (
                    not duracion.is_finite() or duracion != duracion.to_integral_value()
                ):
                    raise ValueError("Duración en meses enteros requerida")
                fecha_revisada = fila["fecha_revision"]
                if isinstance(fecha_revisada, datetime):
                    fecha_revisada = fecha_revisada.date()
                elif fecha_revisada is not None and not isinstance(fecha_revisada, date):
                    fecha_revisada = date.fromisoformat(str(fecha_revisada))
                institucion = texto(fila["institucion"]) or ""
                fuentes = tuple(
                    Fuente(
                        campo,
                        texto(fila.get(referencia)),
                        (texto(fila.get(estado)) or "Sin estado registrado")
                        if estado
                        else "Según PoC; sin verificar",
                    )
                    for campo, referencia, estado in [
                        ("Carrera", "fuente_carrera", None),
                        ("Duración", "fuente_duracion", None),
                        ("Modalidad", "fuente_carrera", "estado_modalidad"),
                        ("Costo", "fuente_costo", "estado_costo"),
                        ("Requisitos", "fuente_requisitos", None),
                        ("Campo laboral", "fuente_campo_laboral", None),
                    ]
                )
                carrera = Carrera(
                    id=texto(fila["id_carrera"]) or "",
                    nombre=texto(fila["carrera"]) or "",
                    institucion=institucion,
                    tipo_formacion="Universitaria" if institucion in UNIVERSIDADES else None,
                    modalidad=texto(fila["modalidad"]),
                    ubicacion=texto(fila["ubicacion"]),
                    duracion_meses=int(duracion) if duracion is not None else None,
                    costo_mensual_soles=numerico("costo_mensual_soles"),
                    requisitos=texto(fila["requisitos"]),
                    campo_laboral=texto(fila["campo_laboral"]),
                    fecha_revision=fecha_revisada,
                    fuentes=fuentes,
                )
                if carrera.id in identificadores or carrera.identidad in identidades:
                    raise ValueError(
                        "Oferta duplicada por código o carrera/institución/sede/modalidad"
                    )
                identificadores.add(carrera.id)
                identidades.add(carrera.identidad)
                resultado.append(carrera)
            except (ValueError, InvalidOperation, TypeError) as error:
                raise ValueError(f"Fila {numero_fila}: {error}") from error
        if not resultado:
            raise ValueError("El libro no contiene ofertas")
        return resultado
    finally:
        libro.close()
